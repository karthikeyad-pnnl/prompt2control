"""Compile Guideline 36 addendum evidence into reports and a workbook."""
from __future__ import annotations

import csv
import difflib
import json
import re
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
EVIDENCE = BASE / "addenda_extracted"
TEMPLATE = BASE / "template.xlsx"
OUT = BASE / "compiled_reports"

SYSTEMS = [
    "Air handling units",
    "Multiple-zone VAV air handling units",
    "Single-zone VAV air handling units",
    "Fan-powered terminal units",
    "VAV terminal units",
    "Dual-duct systems",
    "Chilled-water plants",
    "Hot-water plants",
    "Condenser-water systems",
    "Hydronic distribution systems",
    "Ventilation systems",
    "Economizer systems",
    "Indoor-air-quality and outdoor-air-pollution control systems",
    "Humidity-control systems",
    "Central plants",
    "Generic sequences or cross-system functions",
]


def clean_text(value: str) -> str:
    value = value.replace("", "").replace("�", "")
    value = re.sub(r"\s+", " ", value).strip()
    return value


def clean_description(value: str) -> str:
    value = clean_text(value)
    value = re.sub(r"\s+Description of Change$", "", value, flags=re.I)
    return value


def useful_text(item: dict[str, Any] | None) -> bool:
    if not item or item.get("status") != "found" or not item.get("text"):
        return False
    text = str(item["text"])
    # The extractor can select a contents-page entry for short top-level sections.
    return not ("..." in text or "…" in text or "ASHRAE Guideline 36-2024" in text and len(text) < 180)


def compact_excerpt(text: str | None, limit: int = 650) -> str:
    if not text:
        return "Not located in the supplied edition."
    text = clean_text(text)
    if len(text) <= limit:
        return text
    return text[:limit].rsplit(" ", 1)[0] + " …"


def changed_excerpt(old: str | None, new: str | None, limit: int = 850) -> str:
    if not old and new:
        return "New section/content in 2024; no corresponding 2021 section was extracted."
    if old and not new:
        return "Section/content present in 2021 but not extracted from 2024."
    if not old or not new:
        return "Comparison unavailable from the supplied extraction evidence."
    old_words = clean_text(old).split()
    new_words = clean_text(new).split()
    matcher = difflib.SequenceMatcher(None, old_words, new_words, autojunk=False)
    chunks: list[str] = []
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            continue
        left = " ".join(old_words[max(0, i1 - 18): min(i2 + 18, len(old_words))])
        right = " ".join(new_words[max(0, j1 - 18): min(j2 + 18, len(new_words))])
        chunks.append(f"{tag}: 2021 [{left}] -> 2024 [{right}]")
        if sum(len(c) for c in chunks) > limit:
            break
    if not chunks:
        return "No textual difference detected by normalized extraction; verify formatting and attribution manually."
    result = " | ".join(chunks)
    return result[:limit].rsplit(" ", 1)[0] + (" …" if len(result) > limit else "")


def systems_for(addendum: str, sections: list[str], description: str) -> list[str]:
    joined = " ".join(sections) + " " + description.lower()
    found: list[str] = []
    def add(*items: str) -> None:
        for item in items:
            if item in SYSTEMS and item not in found:
                found.append(item)
    if "humidity" in joined or "dew point" in joined:
        add("Humidity-control systems", "Air handling units", "Multiple-zone VAV air handling units", "VAV terminal units")
    if "ventilation" in joined or "outdoor air" in joined or "outdoor airflow" in joined or "62.1" in joined:
        add("Ventilation systems", "Air handling units", "Multiple-zone VAV air handling units")
    if "economizer" in joined:
        add("Economizer systems", "Air handling units")
    if any(s.startswith("5.16") for s in sections):
        add("Multiple-zone VAV air handling units", "Air handling units")
    if any(s.startswith("5.18") for s in sections):
        add("Single-zone VAV air handling units", "Air handling units")
    if any(s.startswith(x) for s in sections for x in ("5.5", "5.6", "5.7", "5.8", "5.9", "5.10", "5.11", "5.12", "5.13", "5.14")):
        add("VAV terminal units")
    if "fan-powered" in joined:
        add("Fan-powered terminal units")
    if "dual-duct" in joined or any(s.startswith("5.17") for s in sections):
        add("Dual-duct systems")
    if "condenser" in joined or any(s.startswith("5.20") for s in sections):
        add("Condenser-water systems", "Central plants")
    if "chilled" in joined or any(s.startswith("5.19") for s in sections):
        add("Chilled-water plants", "Central plants")
    if "hot water" in joined or "reheat" in joined or any(s.startswith("5.21") for s in sections):
        add("Hot-water plants", "Hydronic distribution systems")
    if "fan" in joined or "pump" in joined or any(s.startswith("5.22") for s in sections):
        add("Air handling units" if "fan" in joined else "Central plants")
    if any(s.startswith("3.") or s.startswith("4.") for s in sections):
        add("Generic sequences or cross-system functions")
    if not found:
        add("Generic sequences or cross-system functions")
    return found


def functions_for(description: str, sections: list[str]) -> tuple[str, list[str]]:
    text = (description + " " + " ".join(sections)).lower()
    categories: list[str] = []
    def add(*items: str) -> None:
        for item in items:
            if item not in categories:
                categories.append(item)
    if any(x in text for x in ("alarm", "alarms")):
        add("Alarm")
    if any(x in text for x in ("humidity", "dew point")):
        add("Humidity control", "Setpoint generation")
    if any(x in text for x in ("ventilation", "outdoor airflow", "outdoor air")):
        add("Ventilation control", "Outdoor-air control")
    if any(x in text for x in ("economizer", "damper")):
        add("Economizer control", "Airflow control")
    if any(x in text for x in ("request", "trim", "respond", "setpoint")):
        add("Setpoint generation")
    if any(x in text for x in ("mode", "occupied", "warm-up", "warmup", "cool-down", "cooldown", "setback", "setup")):
        add("Mode or state logic")
    if any(x in text for x in ("fan", "pump", "enable", "cycle off")):
        add("Equipment enable/disable", "Process control")
    if "afdd" in text or "fault condition" in text:
        add("Fault detection and diagnostics")
    if "condenser water" in text or "plant" in text:
        add("Hydronic control", "Plant optimization")
    if "sensor" in text or "airflow monitoring" in text:
        add("Sensor validation or calibration")
    if "typo" in text or "rearrange" in text or "relocat" in text or "removes the need" in text:
        add("Documentation or editorial correction")
    if not categories:
        add("Process control")
    return categories[0], categories[1:]


def load_inventory() -> dict[str, dict[str, Any]]:
    return json.loads((EVIDENCE / "addenda_inventory.json").read_text(encoding="utf-8"))


def load_comparison(letter: str) -> dict[str, Any]:
    p = EVIDENCE / letter / "section_comparisons.json"
    return json.loads(p.read_text(encoding="utf-8"))


def make_row(letter: str, inventory: dict[str, Any], comparison: dict[str, Any]) -> dict[str, Any]:
    sections = list(inventory["sections_affected"])
    description = clean_description(inventory["description_of_change"])
    section_items = []
    statuses: list[str] = []
    for section in sections:
        old = comparison.get("comparison", {}).get("2021", {}).get(section)
        new = comparison.get("comparison", {}).get("2024", {}).get(section)
        old_ok, new_ok = useful_text(old), useful_text(new)
        if new_ok and (old_ok or old and old.get("status") == "not found"):
            statuses.append("complete")
        else:
            statuses.append("partial")
        section_items.append({
            "section": section,
            "2021": old,
            "2024": new,
            "2021_excerpt": compact_excerpt(old.get("text") if old_ok else None),
            "2024_excerpt": compact_excerpt(new.get("text") if new_ok else None),
            "change_excerpt": changed_excerpt(old.get("text") if old_ok else None, new.get("text") if new_ok else None),
            "2021_pages": old.get("pdf_pages", []) if old else [],
            "2024_pages": new.get("pdf_pages", []) if new else [],
        })
    primary, secondary = functions_for(description, sections)
    systems = systems_for(letter, sections, description)
    complete = all(x == "complete" for x in statuses)
    status = "Complete" if complete else "Partial"
    notes = [
        "The 2024 source is a consolidated edition; section differences cannot be attributed exclusively to this addendum without the controlled standalone addendum redline.",
        "System categories are engineering mappings from affected sections and descriptions; confirm against the official 2024 index.",
    ]
    if not complete:
        notes.append("At least one affected section was not cleanly isolated by the PDF extractor and requires manual page-level review.")
    if letter == "z":
        notes.append("Appendix C description contains a trailing repeated column label; it was removed during normalization.")
    return {
        "addendum": letter,
        "sections_affected": sections,
        "description": description,
        "approval_date": inventory["ashrae_approval_date"],
        "systems_impacted": systems,
        "functional_primary": primary,
        "functional_secondary": secondary,
        "comparison_status": status,
        "quality_level": "Partially verified",
        "appendix_pages": inventory.get("appendix_c_pdf_pages", []),
        "sections": section_items,
        "reviewer_notes": notes,
    }


def report_markdown(row: dict[str, Any]) -> str:
    sections = "; ".join(row["sections_affected"])
    out = [
        f"# ASHRAE Guideline 36-2024 Addendum {row['addendum']} Report",
        "",
        "## Inventory",
        f"- **Description:** {row['description']}",
        f"- **ASHRAE approval date:** {row['approval_date']}",
        f"- **Affected sections:** {sections}",
        f"- **Appendix C PDF page(s):** {', '.join(map(str, row['appendix_pages'])) or 'not recorded'}",
        "",
        "## Methodology and limitations",
        "The supplied `addendum_a_analysis.py` extraction workflow was run in `--all` mode against the supplied 2021 and 2024 PDFs. It extracted Appendix C inventory rows and section-level evidence. The comparisons below use that evidence and normalized word-level alignment; they are not a substitute for the controlled standalone addendum redline. Because the 2024 PDF is consolidated, unrelated cumulative edits in an affected section are not independently attributed to this addendum.",
        "",
        "## Impact classification",
        f"- **Systems impacted:** {'; '.join(row['systems_impacted'])}",
        f"- **Primary functional category:** {row['functional_primary']}",
        f"- **Secondary functional categories:** {', '.join(row['functional_secondary']) or 'None'}",
        f"- **Comparison status:** {row['comparison_status']}",
        f"- **Quality level:** {row['quality_level']}",
        "",
        "## Section-level comparison",
    ]
    for item in row["sections"]:
        out += [
            f"### Section {item['section']}",
            f"- **2021 reference:** PDF page(s) {', '.join(map(str, item['2021_pages'])) or 'not located'}",
            f"- **2024 reference:** PDF page(s) {', '.join(map(str, item['2024_pages'])) or 'not located'}",
            f"- **2021 excerpt:** {item['2021_excerpt']}",
            f"- **2024 excerpt:** {item['2024_excerpt']}",
            f"- **Textual difference:** {item['change_excerpt']}",
            f"- **Operational interpretation:** {row['description']}. The affected behavior should be confirmed against the full section and controlled addendum source before implementation.",
            "",
        ]
    out += ["## Reviewer notes"] + [f"- {note}" for note in row["reviewer_notes"]] + [
        "",
        "## Validation",
        f"- Appendix C inventory row represented: **Pass**.",
        f"- All {len(row['sections_affected'])} listed section references represented: **Pass**.",
        "- Systems and functional categories populated: **Pass with mapping limitation**.",
        f"- 2021-to-2024 evidence status: **{row['comparison_status']}**.",
        "- CDL implementation status: **Not assessed by this PDF-based workflow**.",
        "",
    ]
    return "\n".join(out)


def write_csv(rows: list[dict[str, Any]]) -> None:
    fields = ["Addendum", "Section(s) Affected", "Description of Change", "ASHRAE Approval Date", "Systems Impacted", "Functional Category", "Actual Change from 2021", "2021 Reference", "2024 Reference", "Comparison Status", "Quality Level", "Reviewer Notes"]
    with (OUT / "addenda_impact_matrix.csv").open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            changes = []
            refs21, refs24 = [], []
            for item in row["sections"]:
                changes.append(f"{item['section']}: 2021—{item['2021_excerpt']} | 2024—{item['2024_excerpt']} | Difference—{item['change_excerpt']}")
                if item["2021_pages"]: refs21.append(f"{item['section']}: PDF pp. {', '.join(map(str, item['2021_pages']))}")
                if item["2024_pages"]: refs24.append(f"{item['section']}: PDF pp. {', '.join(map(str, item['2024_pages']))}")
            writer.writerow({
                "Addendum": row["addendum"],
                "Section(s) Affected": "; ".join(row["sections_affected"]),
                "Description of Change": row["description"],
                "ASHRAE Approval Date": row["approval_date"],
                "Systems Impacted": "; ".join(row["systems_impacted"]),
                "Functional Category": f"primary: {row['functional_primary']}; secondary: {', '.join(row['functional_secondary']) or 'None'}",
                "Actual Change from 2021": " || ".join(changes),
                "2021 Reference": "; ".join(refs21) or "Not located",
                "2024 Reference": "; ".join(refs24) or "Not located",
                "Comparison Status": row["comparison_status"],
                "Quality Level": row["quality_level"],
                "Reviewer Notes": " ".join(row["reviewer_notes"]),
            })


def write_workbook(rows: list[dict[str, Any]]) -> None:
    wb = load_workbook(TEMPLATE)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    while ws.max_row > 1:
        ws.delete_rows(2)
    for row in rows:
        changed = " ".join(f"{x['section']}: {x['change_excerpt']}" for x in row["sections"])
        values = {
            "Addendum": row["addendum"],
            "Description": row["description"],
            "Affected sequences": "; ".join(row["sections_affected"]),
            "Affected systems": "; ".join(row["systems_impacted"]),
            "Summary of changed sequences": changed,
            "Implemented in CDL?": "Not assessed (PDF/source comparison only)",
        }
        ws.append([values.get(h, "") for h in headers])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = {1: 11, 2: 58, 3: 38, 4: 52, 5: 100, 6: 32}
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.row_dimensions[1].height = 32
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 105
    wb.save(OUT / "template_filled.xlsx")


def write_yaml_like(rows: list[dict[str, Any]]) -> None:
    payload = {
        "metadata": {
            "title": "ASHRAE Guideline 36-2024 Addenda Impact Matrix",
            "baseline": "ASHRAE Guideline 36-2021",
            "source_workflow": "addendum_a_analysis.py --all",
            "generated_from": "addenda_extracted",
            "quality_note": "All rows are partially verified because the 2024 PDF is consolidated and index confirmation/standalone addendum redlines were not supplied.",
        },
        "rows": rows,
    }
    (OUT / "addenda_impact_matrix.yaml").write_text(
        yaml.safe_dump(payload, sort_keys=False, allow_unicode=True, width=120),
        encoding="utf-8",
    )


def write_validation(rows: list[dict[str, Any]]) -> None:
    partial = [r["addendum"] for r in rows if r["comparison_status"] != "Complete"]
    lines = [
        "# Addenda Compilation Validation Report", "",
        f"- Addenda in Appendix C inventory: **{len(rows)}**.",
        f"- Reports generated: **{len(rows)}**.",
        f"- Workbook rows generated: **{len(rows)}** (excluding the header).",
        "- Original template columns preserved: **Pass**.",
        "- All Appendix C identifiers, descriptions, section lists, and approval dates carried forward: **Pass**.",
        "- System and functional categories populated: **Pass**, using traceable section/description mappings pending official index review.",
        f"- Clean section-level extraction status: **{len(rows) - len(partial)} Complete; {len(partial)} Partial**.",
        f"- Partial extraction addenda: {', '.join(partial) if partial else 'None'}.",
        "- CDL implementation status: **Not assessed**; the supplied workflow and PDFs do not include the CDL implementation repository or a validation run.",
        "",
        "## Required engineering follow-up",
        "1. Review each report against the controlled standalone addendum redline, not only the consolidated 2024 PDF.",
        "2. Confirm system mappings against the official Guideline 36-2024 index.",
        "3. Manually inspect sections marked Partial and any extracted page span that includes a contents or Appendix C page.",
        "4. Assess implementation in the target CDL codebase before marking the workbook's final column Yes/No.",
    ]
    (OUT / "validation_report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    OUT.mkdir(exist_ok=True)
    inventory = load_inventory()
    rows = [make_row(letter, record, load_comparison(letter)) for letter, record in sorted(inventory.items())]
    for row in rows:
        (OUT / f"addendum_{row['addendum']}_report.md").write_text(report_markdown(row), encoding="utf-8")
    write_csv(rows)
    write_yaml_like(rows)
    write_workbook(rows)
    write_validation(rows)
    print(f"Generated {len(rows)} addendum reports in {OUT}")


if __name__ == "__main__":
    main()
